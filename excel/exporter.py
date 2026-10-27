"""Geração da planilha de resultados a partir dos dados persistidos no SQLite.

O arquivo é gravado primeiro em um temporário na mesma pasta e depois substitui o destino
(``os.replace``), de modo que uma interrupção nunca deixa um .xlsx corrompido. Se o arquivo
de destino estiver aberto no Excel, é gerada uma cópia com data/hora no nome.
"""
from __future__ import annotations

import os
import re
import time
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import ITEM_RESULT_SUMMARY, ItemStatus
from database.repository import ItemRecord, Repository, RunInfo
from utils.logger import get_logger

log = get_logger("exporter")

EXPORT_FILTERS: dict[str, tuple[str, tuple[ItemStatus, ...] | None]] = {
    "all": ("Todos os registros", None),
    "found": ("Somente concluídos (encontrados)", (ItemStatus.FOUND,)),
    "not_found": ("Somente sem resultado", (ItemStatus.NOT_FOUND,)),
    "errors": ("Somente erros", (ItemStatus.ERROR,)),
    "processed": ("Consultados (encontrados + sem resultado)", (ItemStatus.FOUND, ItemStatus.NOT_FOUND)),
    "pending": ("Somente pendentes", (ItemStatus.PENDING, ItemStatus.PROCESSING)),
}

COL_RESULT = "RESULTADO"
COL_STATUS = "SITUAÇÃO"
COL_QUERIED_AT = "DATA/HORA DA CONSULTA"
COL_ATTEMPTS = "TENTATIVAS"
COL_MESSAGE = "MENSAGEM"
COL_ROW = "LINHA NA PLANILHA ORIGINAL"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {
    ItemStatus.FOUND.label: PatternFill("solid", fgColor="E2EFDA"),
    ItemStatus.NOT_FOUND.label: PatternFill("solid", fgColor="FFF2CC"),
    ItemStatus.ERROR.label: PatternFill("solid", fgColor="F8CBAD"),
}


def default_output_path(run: RunInfo, output_dir: str | Path) -> Path:
    stem = Path(run.source_file).stem
    return Path(output_dir) / f"{stem}_resultado_exec{run.id}.xlsx"


def _unique(name: str, used: set[str]) -> str:
    candidate, n = name, 1
    while candidate.casefold() in used:
        candidate = f"{name} (consulta)" if n == 1 else f"{name} (consulta {n})"
        n += 1
    used.add(candidate.casefold())
    return candidate


_SMALL_INT = re.compile(r"^-?[1-9]\d{0,8}$|^0$")


def _restore_number(value: str):
    """Colunas que não são a pesquisada: inteiros curtos (até 9 dígitos, sem zero à esquerda) voltam a
    ser números; identificadores longos ou com zeros à esquerda continuam como texto."""
    return int(value) if isinstance(value, str) and _SMALL_INT.match(value) else value


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return str(value)
    return value


def build_dataframe(
    repo: Repository,
    run_id: int,
    export_filter: str = "all",
    field_labels: Sequence[tuple[str, str]] = (),
) -> pd.DataFrame:
    """Monta a tabela de saída: colunas originais + campos capturados + colunas de controle."""
    if export_filter not in EXPORT_FILTERS:
        raise ValueError(f"Filtro de exportação inválido: {export_filter}")
    run = repo.get_run(run_id)
    all_items = repo.items(run_id)
    by_id = {item.id: item for item in all_items}

    statuses = EXPORT_FILTERS[export_filter][1]
    selected: Iterable[ItemRecord] = (
        all_items if statuses is None else (i for i in all_items if i.status in statuses)
    )

    # campos de resultado: os declarados pelo adaptador primeiro, depois quaisquer outros recebidos
    field_keys = [key for key, _ in field_labels]
    labels = dict(field_labels)
    for item in all_items:
        for key in (item.result or {}):
            if key not in labels:
                labels[key] = key
                field_keys.append(key)

    used = {c.casefold() for c in run.columns}
    field_columns = {key: _unique(labels[key], used) for key in field_keys}
    control = {name: _unique(name, used) for name in (COL_RESULT, COL_STATUS, COL_QUERIED_AT, COL_ATTEMPTS, COL_MESSAGE, COL_ROW)}

    rows = []
    for item in selected:
        source = item
        if item.status == ItemStatus.DUPLICATE and item.duplicate_of in by_id:
            source = by_id[item.duplicate_of]  # repete o resultado da primeira ocorrência
        result = source.result or {}
        row = {
            col: item.original.get(col, "") if col == run.search_column else _restore_number(item.original.get(col, ""))
            for col in run.columns
        }
        for key in field_keys:
            row[field_columns[key]] = result.get(key, "")
        summary = ITEM_RESULT_SUMMARY[source.status]
        if item.status == ItemStatus.DUPLICATE:
            summary = f"{ITEM_RESULT_SUMMARY[source.status]} (cópia da linha {source.row_number})"
        row[control[COL_RESULT]] = summary
        row[control[COL_STATUS]] = item.status.label
        row[control[COL_QUERIED_AT]] = datetime.fromisoformat(source.queried_at) if source.queried_at else None
        row[control[COL_ATTEMPTS]] = item.attempts
        row[control[COL_MESSAGE]] = item.message or ""
        row[control[COL_ROW]] = item.row_number
        rows.append(row)

    columns = list(run.columns) + list(field_columns.values()) + list(control.values())
    frame = pd.DataFrame(rows, columns=columns, dtype=object)
    frame.attrs["status_column"] = control[COL_STATUS]
    frame.attrs["datetime_column"] = control[COL_QUERIED_AT]
    return frame


def _write_workbook(frame: pd.DataFrame, summary: list[tuple[str, object]], path: Path) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Resultados")
    status_col = frame.attrs.get("status_column")
    datetime_col = frame.attrs.get("datetime_column")

    for idx, name in enumerate(frame.columns, start=1):
        sample = [len(str(v)) for v in frame[name].head(200).tolist()]
        width = min(max([len(str(name))] + sample) + 2, 60)
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 10)
    ws.freeze_panes = "A2"
    if len(frame.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"

    header = []
    for name in frame.columns:
        cell = WriteOnlyCell(ws, value=_clean(str(name)))
        cell.data_type = "s"
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        header.append(cell)
    ws.append(header)

    for values in frame.itertuples(index=False, name=None):
        out = []
        for name, value in zip(frame.columns, values):
            value = _clean(value)
            cell = WriteOnlyCell(ws, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # nunca interpretar textos iniciados por "=" como fórmula
            if name == datetime_col and value != "":
                cell.number_format = "DD/MM/YYYY HH:MM:SS"
            if name == status_col and value in STATUS_FILLS:
                cell.fill = STATUS_FILLS[value]
            out.append(cell)
        ws.append(out)

    info = wb.create_sheet("Resumo")
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 80
    for key, value in summary:
        key_cell = WriteOnlyCell(info, value=_clean(key))
        key_cell.font = Font(bold=True)
        val_cell = WriteOnlyCell(info, value=_clean(value))
        if isinstance(val_cell.value, str):
            val_cell.data_type = "s"
        info.append([key_cell, val_cell])
    wb.save(path)


def atomic_save(write, target: Path, retries: int = 3, wait: float = 0.5) -> Path:
    """Executa ``write(tmp_path)`` e move o resultado para ``target`` de forma atômica."""
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(f".~{target.stem}.{os.getpid()}.tmp.xlsx")
    try:
        write(tmp)
        with zipfile.ZipFile(tmp) as zf:  # confere a integridade antes de substituir
            bad = zf.testzip()
            if bad:
                raise OSError(f"Arquivo gerado corrompido ({bad})")
        for attempt in range(retries):
            try:
                os.replace(tmp, target)
                return target
            except PermissionError:
                if attempt < retries - 1:
                    time.sleep(wait)
        fallback = target.with_name(f"{target.stem}_{datetime.now():%Y%m%d_%H%M%S}{target.suffix}")
        os.replace(tmp, fallback)
        log.warning("O arquivo %s está aberto/bloqueado; resultado salvo em %s", target.name, fallback.name)
        return fallback
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


def export_run(
    repo: Repository,
    run_id: int,
    output_path: str | Path,
    export_filter: str = "all",
    field_labels: Sequence[tuple[str, str]] = (),
) -> Path:
    run = repo.get_run(run_id)
    frame = build_dataframe(repo, run_id, export_filter, field_labels)
    stats = repo.stats(run_id)
    summary: list[tuple[str, object]] = [
        ("Execução", run.id),
        ("Nome", run.name),
        ("Site/adaptador", run.adapter_id),
        ("Planilha de origem", run.source_file),
        ("Aba / coluna pesquisada", f"{run.sheet} / {run.search_column}"),
        ("Situação da execução", run.status.label),
        ("Filtro da exportação", EXPORT_FILTERS[export_filter][0]),
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Registros consultáveis", stats.total),
        ("Concluídos (encontrados)", stats.found),
        ("Sem resultado", stats.not_found),
        ("Erros", stats.error),
        ("Pendentes", stats.remaining),
        ("Duplicados (não consultados)", stats.duplicate),
        ("Vazios/inválidos (não consultados)", stats.invalid),
        ("Linhas nesta planilha", len(frame)),
    ]
    saved = atomic_save(lambda tmp: _write_workbook(frame, summary, tmp), Path(output_path))
    log.debug("Excel exportado: %s (%d linhas)", saved, len(frame))
    return saved
