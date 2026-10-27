"""Importação de planilhas Excel (.xlsx) para a lista de processamento.

Todas as células são lidas como texto para preservar identificadores (zeros à esquerda,
números longos sem notação científica). Linhas vazias, identificadores em branco,
duplicados e (opcionalmente) inválidos são identificados antes do processamento.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import utils.xlsx_compat  # noqa: F401  (preserva números em colunas formatadas como data)
from core.config import ImportOptions, resolve_path
from core.models import ConfigError, ItemStatus
from utils.validators import cell_to_text, normalize_identifier, validate_identifier

HEADER_SCAN_LIMIT = 50


@dataclass
class SheetInfo:
    name: str
    columns: list[str]
    header_row: int
    data_rows: int


@dataclass
class ImportRecord:
    position: int                 # nº do registro (1 = primeira linha de dados)
    row_number: int               # linha correspondente no Excel
    identifier: str
    original: dict[str, str]
    status: ItemStatus = ItemStatus.PENDING
    message: str = ""
    duplicate_of_row: int | None = None


@dataclass
class ImportResult:
    file: str
    sheet: str
    column: str
    columns: list[str]
    header_row: int
    records: list[ImportRecord] = field(default_factory=list)
    total_data_rows: int = 0      # linhas com dados na aba (antes do intervalo)
    blank_rows: int = 0           # linhas totalmente vazias ignoradas

    def count(self, status: ItemStatus) -> int:
        return sum(1 for r in self.records if r.status == status)

    @property
    def to_process(self) -> int:
        return self.count(ItemStatus.PENDING)

    @property
    def duplicates(self) -> list[ImportRecord]:
        return [r for r in self.records if r.status == ItemStatus.DUPLICATE]

    @property
    def invalid(self) -> list[ImportRecord]:
        return [r for r in self.records if r.status == ItemStatus.INVALID]

    def summary(self) -> str:
        return (
            f"{len(self.records)} registro(s) selecionado(s): {self.to_process} para consultar, "
            f"{len(self.duplicates)} duplicado(s), {len(self.invalid)} vazio(s)/inválido(s)."
        )


def _open_workbook(path: Path):
    if not path.exists():
        raise ConfigError(f"Planilha não encontrada: {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ConfigError("Selecione um arquivo Excel .xlsx.")
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # arquivo corrompido, protegido por senha etc.
        raise ConfigError(f"Não foi possível abrir a planilha: {exc}") from exc


def _read_rows(ws) -> list[list[str]]:
    ws.reset_dimensions()  # não confiar na dimensão gravada no arquivo
    rows: list[list[str]] = []
    for row in ws.iter_rows():
        rows.append([cell_to_text(cell.value, getattr(cell, "number_format", None)) for cell in row])
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _header_names(raw: list[str], width: int) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for idx in range(width):
        name = raw[idx].strip() if idx < len(raw) else ""
        name = name or f"Coluna {get_column_letter(idx + 1)}"
        key = name.casefold()
        if key in seen:
            seen[key] += 1
            name = f"{name} ({seen[key]})"
        else:
            seen[key] = 1
        names.append(name)
    return names


def _sheet_frame(rows: list[list[str]], header_row: int | None) -> tuple[pd.DataFrame, int, int]:
    """Retorna (DataFrame com texto, linha do cabeçalho, linhas vazias ignoradas)."""
    if header_row is None:
        header_row = next((i + 1 for i, r in enumerate(rows[:HEADER_SCAN_LIMIT]) if any(r)), 1)
    if not rows or header_row > len(rows):
        return pd.DataFrame(), header_row, 0

    width = max(len(r) for r in rows)
    # remove colunas finais completamente vazias (inclusive sem título)
    while width > 0 and not any(len(r) >= width and r[width - 1] for r in rows[header_row - 1:]):
        width -= 1
    columns = _header_names(rows[header_row - 1], width)

    data, index, blank = [], [], 0
    for offset, row in enumerate(rows[header_row:], start=header_row + 1):
        values = (row + [""] * width)[:width]
        if not any(values):
            blank += 1
            continue
        data.append(values)
        index.append(offset)
    frame = pd.DataFrame(data, columns=columns, index=pd.Index(index, name="linha_excel"), dtype=object)
    return frame, header_row, blank


def inspect_workbook(path: str | Path) -> list[SheetInfo]:
    """Lista as abas com suas colunas e quantidade de linhas de dados."""
    wb = _open_workbook(Path(path))
    try:
        infos = []
        for ws in wb.worksheets:
            frame, header_row, _ = _sheet_frame(_read_rows(ws), None)
            infos.append(SheetInfo(ws.title, list(frame.columns), header_row, len(frame)))
        return infos
    finally:
        wb.close()


def read_sheet(path: str | Path, sheet: str | None = None, header_row: int | None = None) -> tuple[pd.DataFrame, SheetInfo, int]:
    """Lê uma aba como DataFrame de texto. Índice = número da linha no Excel."""
    wb = _open_workbook(Path(path))
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise ConfigError(f"Aba '{sheet}' não existe. Abas disponíveis: {', '.join(wb.sheetnames)}")
        frame, detected_header, blank = _sheet_frame(_read_rows(ws), header_row)
        return frame, SheetInfo(ws.title, list(frame.columns), detected_header, len(frame)), blank
    finally:
        wb.close()


def build_processing_list(
    frame: pd.DataFrame,
    column: str,
    *,
    start: int | None = None,
    end: int | None = None,
    zero_pad: int | None = None,
    validation: str = "none",
) -> list[ImportRecord]:
    if column not in frame.columns:
        raise ConfigError(f"Coluna '{column}' não encontrada. Colunas: {', '.join(map(str, frame.columns))}")
    total = len(frame)
    first = max(start or 1, 1)
    last = min(end or total, total)
    if total and first > total:
        raise ConfigError(f"O registro inicial ({first}) é maior que o total de registros ({total}).")
    selected = frame.iloc[first - 1:last]

    identifiers = selected[column].map(lambda v: normalize_identifier(str(v), zero_pad))
    duplicated = identifiers.duplicated(keep="first") & identifiers.ne("")
    first_row_of: dict[str, int] = {}

    records: list[ImportRecord] = []
    for position, (row_number, row) in enumerate(selected.iterrows(), start=first):
        identifier = identifiers.loc[row_number]
        record = ImportRecord(
            position=position,
            row_number=int(row_number),
            identifier=identifier,
            original={str(k): str(v) for k, v in row.items()},
        )
        if not identifier:
            record.status, record.message = ItemStatus.INVALID, "Identificador vazio"
        elif duplicated.loc[row_number]:
            record.status = ItemStatus.DUPLICATE
            record.duplicate_of_row = first_row_of[identifier]
            record.message = f"Repetido (mesmo valor da linha {record.duplicate_of_row})"
        else:
            first_row_of[identifier] = int(row_number)
            problem = validate_identifier(identifier, validation)
            if problem:
                record.status, record.message = ItemStatus.INVALID, problem
        records.append(record)
    return records


def import_excel(options: ImportOptions) -> ImportResult:
    options.validate()
    path = resolve_path(options.file)
    frame, info, blank = read_sheet(path, options.sheet, options.header_row)
    records = build_processing_list(
        frame,
        options.column,
        start=options.start,
        end=options.end,
        zero_pad=options.zero_pad,
        validation=options.validation,
    )
    return ImportResult(
        file=str(path),
        sheet=info.name,
        column=options.column,
        columns=info.columns,
        header_row=info.header_row,
        records=records,
        total_data_rows=info.data_rows,
        blank_rows=blank,
    )
