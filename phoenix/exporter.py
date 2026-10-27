"""Exportação: EM_ANDAMENTO, FALHA_PARCIAL_PORTADA, EM_CONFLITO e RELATORIO_TECNICO.

- Cada arquivo recebe só as linhas da sua categoria, em abas com o nome da aba de origem.
- Todas as colunas originais são mantidas; só a coluna "Fila Atual" mapeada é alterada (na cópia).
  Abas sem essa coluna ganham "Fila Atual" ao final, apenas no arquivo exportado.
- Os valores são os calculados pelo Excel (fórmulas não são copiadas, para não quebrar referências).
- Colunas "[ROBÔ] ..." trazem os metadados de auditoria. A gravação é atômica.
"""
from __future__ import annotations

import datetime as dt
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from database.repository import Repository
from excel.exporter import atomic_save
from phoenix.matcher import CATEGORY_LABELS, Category, MatchReport, Outcome, RowResult, match
from phoenix.repository import JobInfo, PhoenixRepository
from utils.logger import get_logger

log = get_logger("phoenix.exporter")

FILES = {
    Category.EM_ANDAMENTO: "EM_ANDAMENTO",
    Category.FALHA_PARCIAL_PORTADA: "FALHA_PARCIAL_PORTADA",
    Category.EM_CONFLITO: "EM_CONFLITO",
}
META = ["[ROBÔ] Aba original", "[ROBÔ] Linha original", "[ROBÔ] CNPJ pesquisado", "[ROBÔ] Número a portar",
        "[ROBÔ] Bilhete", "[ROBÔ] Status do Bilhete", "[ROBÔ] Motivo Recusa", "[ROBÔ] Solicitação",
        "[ROBÔ] Data/hora da pesquisa", "[ROBÔ] Revisão"]
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
META_FILL = PatternFill("solid", fgColor="5B5B5B")
CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")


def _cell(ws, value: Any, **style) -> WriteOnlyCell:
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # nunca virar fórmula
    elif isinstance(value, dt.datetime):
        cell.number_format = "DD/MM/YYYY HH:MM:SS" if (value.hour or value.minute or value.second) else "DD/MM/YYYY"
    elif isinstance(value, dt.date):
        cell.number_format = "DD/MM/YYYY"
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def _sheet_title(name: str, used: set[str]) -> str:
    title = "".join("_" if c in '[]:*?/\\' else c for c in name)[:31] or "Aba"
    base, n = title, 2
    while title.casefold() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.casefold())
    return title


def _solicitacao_text(res: RowResult) -> str:
    if not res.ticket:
        return ""
    dados = res.ticket.solicitacao
    parts = [res.ticket.solicitacao_id] + [f"{k}: {v}" for k, v in dados.items()
                                           if k.lower() not in ("razão social", "razao social", "cnpj") and v]
    return " | ".join(parts)


def _meta(res: RowResult, queried_at: str | None) -> list[Any]:
    t = res.ticket
    return [res.row.sheet, res.row.row, res.row.cnpj.value, t.fields.get("numero_portar", "") if t else "",
            t.fields.get("bilhete", "") if t else "", t.status if t else "", t.motivo if t else "",
            _solicitacao_text(res), queried_at or "", " ".join(res.revisao)]


def _write_category(path: Path, category: Category, results: list[RowResult], job: JobInfo,
                    queried: dict[str, str | None]) -> Path:
    by_sheet: dict[str, list[RowResult]] = defaultdict(list)
    for res in results:
        by_sheet[res.row.sheet].append(res)

    def write(tmp: Path) -> None:
        wb = Workbook(write_only=True)
        used: set[str] = set()
        if not results:
            ws = wb.create_sheet("Sem registros")
            ws.append([f"Nenhuma linha classificada como '{CATEGORY_LABELS[category]}' nesta execução."])
        for sheet in job.sheets:
            name = sheet["mapping"]["sheet"]
            rows = by_sheet.get(name)
            if not rows:
                continue
            headers = list(sheet["headers"])
            fila_letter = sheet["mapping"].get("fila_atual")
            fila_idx = column_index_from_string(fila_letter) - 1 if fila_letter else len(headers)
            if not fila_letter:
                headers.append("Fila Atual")
            ws = wb.create_sheet(_sheet_title(name, used))
            ws.freeze_panes = "A2"
            for j, title in enumerate(headers + META, start=1):
                ws.column_dimensions[get_column_letter(j)].width = 18 if j <= len(headers) else 24
            ws.append([_cell(ws, h, font=HEADER_FONT, fill=HEADER_FILL) for h in headers]
                      + [_cell(ws, h, font=HEADER_FONT, fill=META_FILL) for h in META])
            for res in sorted(rows, key=lambda r: r.row.row):
                values = list(res.row.values) + [None] * (len(headers) - len(res.row.values))
                values = values[:len(headers)]
                values[fila_idx] = res.fila_nova
                cells = [_cell(ws, v, fill=CHANGED_FILL) if j == fila_idx else _cell(ws, v) for j, v in enumerate(values)]
                ws.append(cells + [_cell(ws, v) for v in _meta(res, queried.get(res.row.cnpj.value))])
        wb.save(tmp)

    return atomic_save(write, path)


def _table(wb, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.freeze_panes = "A2"
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    ws.append([_cell(ws, h, font=HEADER_FONT, fill=HEADER_FILL) for h in header])
    for row in rows:
        ws.append([_cell(ws, v) for v in row])
    if not rows:
        ws.append([_cell(ws, "(nenhum registro)")])


def _write_report(path: Path, report: MatchReport, job: JobInfo, prepo: PhoenixRepository, files: dict[str, Path]) -> Path:
    repo = prepo.repo
    run = repo.get_run(job.run_id)
    stats = repo.stats(job.run_id)
    items = repo.items(job.run_id)
    numbers_by_cnpj: dict[str, list[str]] = defaultdict(list)
    for item in items:
        for sol in (item.result or {}).get("solicitacoes", []):
            for t in sol.get("bilhetes", []):
                numbers_by_cnpj[item.identifier].append(t.get("numero_portar", ""))

    def row_line(res: RowResult, extra: str = "") -> list[Any]:
        t = res.ticket
        return [res.row.sheet, res.row.row, res.row.cnpj.raw, res.row.cnpj.value, res.row.acesso_raw, res.outcome.value,
                CATEGORY_LABELS.get(res.category, "") if res.category else "", t.status if t else "",
                t.motivo if t else "", " ".join(res.revisao) or extra]

    row_header = ["Aba", "Linha", "CNPJ (planilha)", "CNPJ (normalizado)", "Acesso (planilha)", "Resultado",
                  "Categoria", "Status do Bilhete", "Motivo Recusa", "Observação"]

    def write(tmp: Path) -> None:
        wb = Workbook(write_only=True)
        summary = report.summary()
        resumo = [
            ["Execução", job.run_id], ["Situação", run.status.label], ["Planilha de origem", job.source_file],
            ["SHA-256 da origem", job.source_sha256], ["Gerado em", dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")],
            ["CNPJs distintos", stats.total], ["CNPJs pesquisados", stats.done], ["CNPJs pendentes", stats.remaining],
        ] + [[k.replace("_", " ").capitalize(), v] for k, v in summary.items()] + [
            ["Arquivo " + name, str(p)] for name, p in files.items()
        ] + [["Regra para vários bilhetes", job.settings.match.multiplos_bilhetes],
             ["Status 'Falha parcial / portada' considerados", "; ".join(job.settings.status.falha_parcial_portada)],
             ["Observação", "Fila Atual é alterada somente nas cópias exportadas; o arquivo de origem não é modificado. "
                            "Os valores exportados são os calculados (fórmulas não são copiadas)."]]
        _table(wb, "Resumo", ["Item", "Valor"], resumo)
        revisar = [res for res in report.rows if res.outcome == Outcome.REVISAR
                   or (res.outcome == Outcome.CLASSIFICADO and res.revisao)]
        _table(wb, "Revisar", row_header, [row_line(r) for r in revisar])
        _table(wb, "Sem correspondência", row_header + ["Números a portar no Phoenix para o CNPJ"],
               [row_line(r) + ["; ".join(n for n in numbers_by_cnpj.get(r.row.cnpj.value, []) if n)]
                for r in report.rows if r.outcome == Outcome.SEM_CORRESPONDENCIA])
        _table(wb, "Sem número ou CNPJ inválido", row_header,
               [row_line(r) for r in report.rows if r.outcome in (Outcome.SEM_NUMERO, Outcome.CNPJ_INVALIDO)])
        _table(wb, "CNPJ sem solicitações", ["CNPJ", "Linhas", "Abas", "Consultado em"],
               [[i.identifier, i.original.get("linhas"), i.original.get("abas"), i.queried_at]
                for i in items if i.status.value == "not_found"])
        _table(wb, "Erros de consulta", ["CNPJ", "Tentativas", "Mensagem", "Abas"],
               [[i.identifier, i.attempts, i.message, i.original.get("abas")] for i in items if i.status.value == "error"])
        _table(wb, "Bilhetes sem linha", ["CNPJ", "Solicitação", "Número a portar", "Bilhete", "Status do Bilhete", "Motivo Recusa"],
               [[t.cnpj, t.solicitacao_id, t.fields.get("numero_portar", ""), t.fields.get("bilhete", ""), t.status, t.motivo]
                for t in report.unmatched_tickets])
        outcome_counts = Counter(r.outcome.value for r in report.rows)
        _table(wb, "Contagens", ["Tipo", "Valor", "Quantidade"],
               [["Status do Bilhete (todos os bilhetes lidos)", k, v] for k, v in report.ticket_status_counts.most_common()]
               + [["Resultado por linha da planilha", k, v] for k, v in outcome_counts.most_common()])
        _table(wb, "Consultas", ["CNPJ", "Situação", "Solicitações", "Bilhetes", "Consultado em", "Tentativas", "Mensagem"],
               [[i.identifier, i.status.label, len((i.result or {}).get("solicitacoes", [])),
                 sum(len(s.get("bilhetes", [])) for s in (i.result or {}).get("solicitacoes", [])),
                 i.queried_at, i.attempts, i.message] for i in items])
        _table(wb, "Mapeamento", ["Aba", "Linha cabeçalho", "CNPJ", "Acesso", "Fila Atual", "Nº temporário", "Observações"],
               [[s["mapping"]["sheet"], s["mapping"].get("header_row"), s["mapping"].get("cnpj"), s["mapping"].get("acesso"),
                 s["mapping"].get("fila_atual") or "(criada na exportação)", s["mapping"].get("numero_temporario"),
                 " ".join([s["mapping"].get("note") or ""] + s.get("issues", []))] for s in job.sheets])
        wb.save(tmp)

    return atomic_save(write, path)


def build_report(prepo: PhoenixRepository, run_id: int) -> tuple[JobInfo, MatchReport]:
    job = prepo.job(run_id)
    report = match(prepo.rows(run_id), prepo.queries(run_id), job.settings.status, job.settings.match)
    return job, report


def export_job(repo: Repository, run_id: int, output_dir: Path) -> dict[str, Path]:
    """Gera (ou regrava) os 3 arquivos e o relatório técnico da execução."""
    prepo = PhoenixRepository(repo)
    job, report = build_report(prepo, run_id)
    stamp = prepo.export_stamp(run_id)
    output_dir.mkdir(parents=True, exist_ok=True)
    queried = {cnpj: q.queried_at for cnpj, q in prepo.queries(run_id).items()}
    files: dict[str, Path] = {}
    for category, prefix in FILES.items():
        files[prefix] = _write_category(output_dir / f"{prefix}_{stamp}.xlsx", category, report.by_category(category),
                                        job, queried)
    files["RELATORIO_TECNICO"] = _write_report(output_dir / f"RELATORIO_TECNICO_{stamp}.xlsx", report, job, prepo, dict(files))
    log.info("Arquivos exportados em %s: %s", output_dir, ", ".join(p.name for p in files.values()))
    return files
