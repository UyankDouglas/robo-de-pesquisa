"""Leitura da planilha de origem (somente leitura): abas, cabeçalhos, mapeamento e validação.

O arquivo de origem nunca é alterado. As linhas são lidas com os valores calculados (fórmulas
viram o último valor salvo pelo Excel) e guardadas para a exportação.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

import utils.xlsx_compat  # noqa: F401  (preserva números em colunas formatadas como data)
from core.models import ConfigError
from phoenix.config import SheetMapping
from phoenix.normalize import CnpjInfo, PhoneRules, cell_text, normalize_cnpj, normalize_label, normalize_phone

HEADER_SCAN = 15
CNPJ_HEADERS = {"cnpj", "cnpj/cpf", "cnpj cpf", "cpf/cnpj", "cnpj do cliente"}
ACESSO_HEADERS = {"acesso", "acesso 1", "acesso1", "numero de acesso"}
FILA_HEADERS = {"fila atual"}
TEMP_HEADERS = {"temporario", "numero temporario", "n temporario"}


@dataclass
class SheetSummary:
    name: str
    hidden: bool
    header_row: int | None
    data_rows: int
    headers: dict[str, str]           # letra -> cabeçalho


@dataclass
class SourceRow:
    sheet: str
    row: int
    cnpj: CnpjInfo
    acesso_raw: str
    acesso: str
    temp_raw: str
    temp: str
    fila_raw: str
    values: list[Any]
    warnings: list[str] = field(default_factory=list)


@dataclass
class SheetData:
    mapping: SheetMapping
    headers: list[str]
    rows: list[SourceRow]
    issues: list[str] = field(default_factory=list)

    @property
    def name(self) -> str:
        return self.mapping.sheet


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _open(path: str | Path):
    path = Path(path)
    if not path.exists():
        raise ConfigError(f"Planilha não encontrada: {path}")
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ConfigError(f"Não foi possível abrir a planilha: {exc}") from exc


def _plain(value: Any) -> Any:
    """Valor para exportar: inteiros sem '.0', datas preservadas, demais como estão."""
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return int(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time, int, float, bool)) or value is None:
        return value
    return str(value)


def _rows(ws) -> list[tuple]:
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    while rows and not any(c not in (None, "") for c in rows[-1]):
        rows.pop()
    return rows


def _detect_header(rows: list[tuple]) -> int | None:
    for i, row in enumerate(rows[:HEADER_SCAN], start=1):
        texts = [c for c in row if isinstance(c, str) and c.strip() and not re.fullmatch(r"[\d\W]+", c.strip())]
        if len(texts) >= 5:
            return i
    return None


def _headers(rows: list[tuple], header_row: int) -> dict[str, str]:
    row = rows[header_row - 1] if 0 < header_row <= len(rows) else ()
    return {get_column_letter(j + 1): cell_text(c) for j, c in enumerate(row) if cell_text(c)}


def list_sheets(path: str | Path) -> list[SheetSummary]:
    wb = _open(path)
    try:
        out = []
        for ws in wb.worksheets:
            rows = _rows(ws)
            header = _detect_header(rows)
            data = sum(1 for r in rows[header or 0:] if any(c not in (None, "") for c in r))
            out.append(SheetSummary(ws.title, ws.sheet_state != "visible", header, data,
                                    _headers(rows, header) if header else {}))
        return out
    finally:
        wb.close()


def _column_values(rows: list[tuple], header_row: int, letter: str) -> list[Any]:
    j = column_index_from_string(letter) - 1
    return [r[j] if j < len(r) else None for r in rows[header_row:] if any(c not in (None, "") for c in r)]


def _share(values: Iterable[Any], predicate) -> float:
    values = [v for v in values if v not in (None, "")]
    return sum(1 for v in values if predicate(v)) / len(values) if values else 0.0


def _is_text(v: Any) -> bool:
    return isinstance(v, str) and not re.fullmatch(r"[\d\s.,/-]+", v.strip())


def detect_mapping(name: str, rows: list[tuple], header_row: int | None, rules: PhoneRules = PhoneRules()) -> SheetMapping:
    """Sugere o mapeamento pelo cabeçalho E pelo conteúdo (nunca só pela 1ª/última ocorrência do nome)."""
    mapping = SheetMapping(sheet=name, header_row=header_row or 1)
    if not header_row:
        mapping.note = "Cabeçalho não identificado: configure manualmente."
        return mapping
    headers = _headers(rows, header_row)
    by_name: dict[str, list[str]] = {}
    for letter, text in headers.items():
        by_name.setdefault(normalize_label(text), []).append(letter)
    notes = []

    def pick(names: set[str], check, label: str) -> str | None:
        letters = [l for n in names for l in by_name.get(n, [])]
        good = [l for l in letters if _share(_column_values(rows, header_row, l), check) >= 0.5]
        if len(letters) > 1:
            notes.append(f"{label}: cabeçalho repetido em {', '.join(letters)}; conteúdo válido em {', '.join(good) or 'nenhuma'}")
        return good[0] if len(good) == 1 else None

    mapping.cnpj = pick(CNPJ_HEADERS, lambda v: normalize_cnpj(v).status == "valido", "CNPJ")
    mapping.acesso = pick(ACESSO_HEADERS, lambda v: bool(normalize_phone(v, rules)), "Acesso")
    mapping.fila_atual = pick(FILA_HEADERS, _is_text, "Fila Atual")
    mapping.numero_temporario = pick(TEMP_HEADERS, lambda v: bool(normalize_phone(v, rules)), "Número temporário")
    if not mapping.acesso:
        notes.append("Sem coluna de número de acesso com telefones: as linhas não podem ser relacionadas.")
    if not mapping.fila_atual and not by_name.get("fila atual"):
        notes.append("Sem coluna 'Fila Atual': será criada somente na planilha exportada.")
    for attr in ("cnpj", "acesso", "fila_atual", "numero_temporario"):
        letter = getattr(mapping, attr)
        if letter:
            mapping.expected_headers[letter] = headers[letter]
    mapping.note = " ".join(notes)
    return mapping


def detect_mappings(path: str | Path, sheets: list[str] | None = None, rules: PhoneRules = PhoneRules()) -> dict[str, SheetMapping]:
    wb = _open(path)
    try:
        out = {}
        for ws in wb.worksheets:
            if sheets is None or ws.title in sheets:
                rows = _rows(ws)
                out[ws.title] = detect_mapping(ws.title, rows, _detect_header(rows), rules)
        return out
    finally:
        wb.close()


def _check_mapping(mapping: SheetMapping, rows: list[tuple]) -> list[str]:
    """Erros que impedem o processamento (coluna inexistente ou cabeçalho diferente do esperado)."""
    errors = []
    if mapping.header_row < 1 or mapping.header_row > len(rows):
        return [f"Linha de cabeçalho {mapping.header_row} fora da aba."]
    headers = _headers(rows, mapping.header_row)
    if not mapping.cnpj:
        errors.append("Coluna CNPJ não configurada.")
    for attr, label in (("cnpj", "CNPJ"), ("acesso", "Acesso"), ("fila_atual", "Fila Atual"), ("numero_temporario", "Nº temporário")):
        letter = getattr(mapping, attr)
        if not letter:
            continue
        try:
            column_index_from_string(letter)
        except ValueError:
            errors.append(f"{label}: coluna '{letter}' inválida.")
            continue
        expected = mapping.expected_headers.get(letter)
        found = headers.get(letter, "")
        if expected is not None and normalize_label(expected) != normalize_label(found):
            errors.append(f"{label}: na coluna {letter} era esperado o cabeçalho '{expected}' e há '{found or '(vazio)'}'.")
    return errors


def load_sheet(ws, mapping: SheetMapping, rules: PhoneRules = PhoneRules()) -> SheetData:
    rows = _rows(ws)
    errors = _check_mapping(mapping, rows)
    if errors:
        raise ConfigError(f"Aba '{mapping.sheet}': " + " ".join(errors))
    header_values = rows[mapping.header_row - 1]
    width = max([len(r) for r in rows[mapping.header_row - 1:]] + [0])
    headers = [cell_text(header_values[j]) if j < len(header_values) else "" for j in range(width)]
    while headers and not headers[-1] and all(len(r) < len(headers) or r[len(headers) - 1] in (None, "") for r in rows):
        headers.pop()
    width = len(headers)

    def idx(letter: str | None) -> int | None:
        return column_index_from_string(letter) - 1 if letter else None

    i_cnpj, i_acc, i_fila, i_temp = (idx(mapping.cnpj), idx(mapping.acesso), idx(mapping.fila_atual), idx(mapping.numero_temporario))
    fila_name = normalize_label(headers[i_fila]) if i_fila is not None and i_fila < width else None
    twins = [j for j, h in enumerate(headers) if fila_name and normalize_label(h) == fila_name and j != i_fila]

    def get(row: tuple, j: int | None) -> Any:
        return row[j] if j is not None and j < len(row) else None

    data = SheetData(mapping, headers, [])
    if not mapping.acesso:
        data.issues.append("Sem coluna de número de acesso: nenhuma linha desta aba pode ser relacionada a um bilhete.")
    for number, row in enumerate(rows[mapping.header_row:], start=mapping.header_row + 1):
        if not any(c not in (None, "") for c in row):
            continue
        acc_raw = cell_text(get(row, i_acc))
        temp_raw = cell_text(get(row, i_temp))
        src = SourceRow(
            sheet=mapping.sheet, row=number, cnpj=normalize_cnpj(get(row, i_cnpj)),
            acesso_raw=acc_raw, acesso=normalize_phone(acc_raw, rules),
            temp_raw=temp_raw, temp=normalize_phone(temp_raw, rules),
            fila_raw=cell_text(get(row, i_fila)),
            values=[_plain(row[j]) if j < len(row) else None for j in range(width)],
        )
        if i_fila is not None and not src.fila_raw and any(_is_text(get(row, t)) for t in twins):
            others = ", ".join(get_column_letter(t + 1) for t in twins if _is_text(get(row, t)))
            src.warnings.append(f"Fila Atual vazia em {mapping.fila_atual} e com texto em {others}: possível "
                                f"deslocamento de colunas nesta linha; revisar.")
        data.rows.append(src)

    shifted = sum(1 for r in data.rows if r.warnings)
    if shifted:
        data.issues.append(f"{shifted} linha(s) com possível deslocamento da coluna Fila Atual (ver relatório técnico).")
    if i_fila is not None:
        numeric = _share((r.values[i_fila] for r in data.rows if i_fila < len(r.values)), lambda v: isinstance(v, (int, float)))
        if numeric > 0.5:
            data.issues.append(f"A coluna Fila Atual ({mapping.fila_atual}) tem maioria de valores numéricos: confira o mapeamento.")
    counts = Counter(r.cnpj.status for r in data.rows)
    if counts.get("invalido") or counts.get("cpf"):
        data.issues.append(f"{counts.get('invalido', 0)} CNPJ(s) inválido(s) e {counts.get('cpf', 0)} CPF(s) não serão consultados.")
    no_phone = sum(1 for r in data.rows if mapping.acesso and not r.acesso)
    if no_phone:
        data.issues.append(f"{no_phone} linha(s) sem número de acesso válido.")
    return data


def load_sheets(path: str | Path, mappings: dict[str, SheetMapping], rules: PhoneRules = PhoneRules()) -> list[SheetData]:
    wb = _open(path)
    try:
        missing = [name for name in mappings if name not in wb.sheetnames]
        if missing:
            raise ConfigError(f"Aba(s) inexistente(s) na planilha: {', '.join(missing)}")
        return [load_sheet(wb[name], mapping, rules) for name, mapping in mappings.items()]
    finally:
        wb.close()
