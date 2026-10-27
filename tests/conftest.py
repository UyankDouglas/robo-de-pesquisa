"""Fixtures compartilhadas pelos testes."""
from __future__ import annotations

import logging
import sys
import threading
import time
from pathlib import Path
from typing import Any, Callable

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from adapters.base_adapter import BaseAdapter  # noqa: E402
from core import paths  # noqa: E402
from core.config import RunSettings  # noqa: E402
from core.models import SearchResult  # noqa: E402
from database.connection import Database  # noqa: E402
from database.repository import Repository  # noqa: E402
from excel.importer import ImportResult, build_processing_list, read_sheet  # noqa: E402


@pytest.fixture(autouse=True)
def _reset_app_logging():
    """Remove handlers criados por setup_logging (evita gravar em streams já fechados pelo pytest)."""
    yield
    logger = logging.getLogger("wsb")
    for handler in list(logger.handlers):
        if getattr(handler, "_wsb_managed", False):
            logger.removeHandler(handler)
            handler.close()
    logger.propagate = True


@pytest.fixture
def tmp_home(tmp_path, monkeypatch) -> Path:
    """Redireciona todas as pastas de dados do aplicativo para um diretório temporário."""
    home = tmp_path / "home"
    mapping = {
        "DATA_ROOT": home,
        "CONFIG_DIR": home / "config",
        "SITES_DIR": home / "config" / "sites",
        "INPUT_DIR": home / "input",
        "OUTPUT_DIR": home / "output",
        "LOGS_DIR": home / "logs",
        "DATA_DIR": home / "data",
        "SESSIONS_DIR": home / "data" / "sessions",
        "LOCKS_DIR": home / "data" / "locks",
        "DB_PATH": home / "data" / "test.db",
        "UI_STATE_PATH": home / "data" / "ui_state.json",
    }
    for name, value in mapping.items():
        monkeypatch.setattr(paths, name, value)
    paths.ensure_dirs()
    return home


@pytest.fixture
def db(tmp_home) -> Database:
    return Database(paths.DB_PATH)


@pytest.fixture
def repo(db):
    repository = Repository(db.connect())
    yield repository
    repository.close()


def make_workbook(path: Path, rows: list[list[Any]], sheet: str = "Dados", formats: dict[str, str] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # texto literal, não fórmula
    for coord, fmt in (formats or {}).items():
        ws[coord].number_format = fmt
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def make_import(tmp_path: Path, identifiers: list[str], column: str = "DOC") -> ImportResult:
    """Cria uma planilha com os identificadores e devolve o resultado da importação."""
    rows = [["ID", column]] + [[i + 1, value] for i, value in enumerate(identifiers)]
    path = make_workbook(tmp_path / f"entrada_{len(identifiers)}_{time.monotonic_ns()}.xlsx", rows)
    frame, info, blank = read_sheet(path)
    records = build_processing_list(frame, column)
    return ImportResult(str(path), info.name, column, info.columns, info.header_row, records, len(frame), blank)


def fast_settings(**overrides) -> RunSettings:
    base = dict(adapter_id="fake", interval_seconds=0, timeout_seconds=5, max_retries=2,
                backoff_base_seconds=0.01, backoff_max_seconds=0.05, autosave_every=1000, autosave_seconds=3600)
    base.update(overrides)
    return RunSettings(**base)


class FakeAdapter(BaseAdapter):
    """Adaptador sem navegador com respostas programáveis por identificador.

    ``script[identifier]`` é uma lista consumida a cada chamada: SearchResult ou exceção.
    Sem script, devolve "encontrado".
    """

    id = "fake"
    display_name = "Fake"
    requires_browser = False
    fields = [("nome", "NOME")]

    def __init__(self, script: dict[str, list[Any]] | None = None, delay: float = 0.0,
                 login_required: bool = False, on_search: Callable[[str], None] | None = None):
        super().__init__({})
        self.script = {k: list(v) for k, v in (script or {}).items()}
        self.delay = delay
        self.calls: list[tuple[str, float]] = []
        self.login_required = login_required
        self.logged_in = not login_required
        self.login_prompts = 0
        self.on_search = on_search
        self.lock = threading.Lock()

    @property
    def requires_login(self) -> bool:
        return self.login_required

    def search(self, identifier: str) -> SearchResult:
        with self.lock:
            self.calls.append((identifier, time.monotonic()))
        if self.on_search:
            self.on_search(identifier)
        if self.delay:
            time.sleep(self.delay)
        queue = self.script.get(identifier)
        outcome = queue.pop(0) if queue else SearchResult.found_with({"nome": f"Nome {identifier}"})
        if isinstance(outcome, BaseException):
            raise outcome
        return outcome

    def is_logged_in(self) -> bool:
        return self.logged_in

    def begin_login(self) -> str:
        self.login_prompts += 1
        return "Faça login (teste)"

    def called(self) -> list[str]:
        return [c[0] for c in self.calls]


@pytest.fixture
def demo_server():
    from demo_site.server import DemoOptions, DemoServer

    server = DemoServer("127.0.0.1", 0, DemoOptions(latency=0.01, slow_seconds=3)).start()
    yield server
    server.stop()


@pytest.fixture
def login_server():
    from demo_site.server import DemoOptions, DemoServer

    server = DemoServer("127.0.0.1", 0, DemoOptions(require_login=True, latency=0.01)).start()
    yield server
    server.stop()


def wait_until(predicate: Callable[[], bool], timeout: float = 10.0, step: float = 0.02) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return True
        time.sleep(step)
    return predicate()
