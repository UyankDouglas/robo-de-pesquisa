"""Dados 100% fictícios para os testes do Phoenix Bot (CNPJs com dígitos verificadores válidos)."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from phoenix.normalize import _cnpj_check


def cnpj(base12: str) -> str:
    return base12 + _cnpj_check(base12)


C_SINGLE = cnpj("100000000001")   # 1 solicitação, 1 bilhete "Em andamento"
C_MULTI = cnpj("200000000001")    # 3 solicitações
C_MIXED = cnpj("300000000001")    # 1 solicitação com bilhetes em status diferentes
C_NONE = cnpj("400000000001")     # "Nenhuma Solicitação foi encontrada."
C_PAGED = cnpj("500000000001")    # 12 solicitações -> 3 páginas
C_ZEROS = cnpj("000123450001")    # zeros à esquerda (o Excel guarda como número)
C_HIST = cnpj("600000000001")     # mesmo número em várias solicitações


def t(numero: str, status: str, motivo: str = "-", bilhete: str = "", temp: str = "") -> dict[str, Any]:
    return {"numero": numero, "status": status, "motivo": motivo, "bilhete": bilhete or f"B{numero[-6:]}", "temp": temp}


def sol(razao: str, data: str, bilhetes: list, status: str = "Concluída") -> dict[str, Any]:
    return {"razao": razao, "data": data, "tipo": "Portabilidade", "status": status, "port_in": "Normal", "bilhetes": bilhetes}


DATA: dict[str, list[dict[str, Any]]] = {
    C_SINGLE: [sol("ALFA FICTICIA LTDA", "01/09/2026 10:00", [t("81999990001", "Em andamento")])],
    C_MULTI: [
        # o "Status Solicitação" (Em andamento) NÃO pode influenciar a classificação
        sol("BETA FICTICIA SA", "02/09/2026 09:00", [t("81999990011", "Portado")], status="Em andamento"),
        sol("BETA FICTICIA SA", "05/09/2026 09:00", [t("81999990012", "Em conflito", "Linha com débito na doadora")]),
        sol("BETA FICTICIA SA", "07/09/2026 09:00", [t("81999990013", "Falha parcial / portada")]),
    ],
    C_MIXED: [sol("GAMA FICTICIA ME", "03/09/2026 14:30", [
        t("85999990021", "Em andamento"),
        t("85999990022", "Falha parcial / portada"),
        t("85999990023", "Em conflito", "Informação disponibilizada pelo sistema."),
        t("85999990024", "Em conflito", "-"),
        t("85999990025", "Cancelado"),
        t("85999990026", "Portado"),
    ])],
    C_PAGED: [sol("DELTA FICTICIA LTDA", f"{d:02d}/08/2026 08:00", [t(f"119999900{d:02d}", "Em andamento" if d == 12 else "Portado")])
              for d in range(1, 13)],
    C_ZEROS: [sol("EPSILON ZEROS LTDA", "04/09/2026 11:00", [t("71999990031", "Em conflito", "Titular divergente")])],
    C_HIST: [
        sol("ZETA HISTORICO LTDA", "01/08/2026 10:00", [t("31999990041", "Em conflito", "Primeira tentativa")]),
        sol("ZETA HISTORICO LTDA", "20/08/2026 10:00", [t("31999990041", "Em andamento"), t("31999990042", "Em andamento")]),
        sol("ZETA HISTORICO LTDA", "20/08/2026 10:00", [t("31999990042", "Em conflito", "Outro motivo")]),
    ],
}

def client_result(cnpj_: str) -> dict[str, Any] | None:
    """O que o PhoenixClient devolveria para o CNPJ (mesmo formato), sem abrir navegador."""
    from phoenix_simulator.server import fmt_phone

    sols = DATA.get(cnpj_)
    if not sols:
        return None
    out = []
    for i, s in enumerate(sols):
        out.append({
            "id": f"pag{i // 5 + 1}-sol{i % 5 + 1}", "pagina": i // 5 + 1, "posicao": i % 5,
            "dados": {"Razão Social": s["razao"], "CNPJ": cnpj_, "Data Solicitação": s["data"], "Tipo": s["tipo"],
                      "Status Solicitação": s["status"], "Tipo Port In": s["port_in"]},
            "bilhetes": [{"numero_portar": fmt_phone(b["numero"]), "bilhete": b["bilhete"],
                          "numero_temporario": fmt_phone(b.get("temp", "")), "status_bilhete": b["status"],
                          "autorizacao_doadora": "Sim", "autorizado_ea": "Sim", "motivo_recusa": b["motivo"],
                          "status_linha": "Ativa", "motivo_falha": "-"} for b in s["bilhetes"]],
        })
    return {"cnpj": cnpj_, "consultado_em": "2026-10-27T10:00:00", "solicitacoes": out}


TLV_HEADERS =["RADAR", "RAZÃO SOCIAL", "CNPJ", "Acesso", "fila atual", "fila atual", "OBS"]
CSTV_HEADERS = ["ID Produto", "Acesso", "Temporário", "Cliente", "CNPJ/CPF", "Status P2B"]


def build_workbook(path: Path) -> Path:
    """Planilha de entrada fictícia com as armadilhas da planilha real (sem dados reais)."""
    wb = Workbook()
    tlv = wb.active
    tlv.title = "TLV"
    tlv.append(TLV_HEADERS)  # C=CNPJ, D=Acesso, E=fila atual (texto), F=fila atual (numérico)
    rows = [
        ("R1", "ALFA", int(C_SINGLE), 81999990001, "CONCLUÍDO", 0, "linha 2"),
        ("R2", "BETA", int(C_MULTI), 81999990011, "CONCLUÍDO", 1, "outro status"),
        ("R3", "BETA", int(C_MULTI), 81999990012, "CONCLUÍDO", 2, "conflito"),
        ("R4", "BETA", int(C_MULTI), 81999990013, "ENTREGA", 3, "falha parcial"),
        ("R5", "GAMA", int(C_MIXED), "(85) 99999-0021", "CONCLUÍDO", 0, "formatado"),
        ("R6", "GAMA", int(C_MIXED), 85999990022, "CONCLUÍDO", 0, ""),
        ("R7", "GAMA", int(C_MIXED), 85999990023, "CONCLUÍDO", 0, ""),
        ("R8", "GAMA", int(C_MIXED), 85999990024, "CONCLUÍDO", 0, "conflito sem motivo"),
        ("R9", "GAMA", int(C_MIXED), 85999990025, "CONCLUÍDO", 0, "cancelado"),
        ("R10", "GAMA", int(C_MIXED), 85999990099, "CONCLUÍDO", 0, "número não existe no Phoenix"),
        ("R11", "INEX", int(C_NONE), 21999990000, "CONCLUÍDO", 0, "cnpj sem solicitações"),
        ("R12", "DELTA", int(C_PAGED), 11999990012, "CONCLUÍDO", 0, "página 3"),
        ("R13", "ZEROS", int(C_ZEROS), 5571999990031, "CONCLUÍDO", 0, "DDI 55"),
        ("R14", "ZETA", int(C_HIST), 31999990041, "CONCLUÍDO", 0, "histórico: vale o mais recente"),
        ("R15", "ZETA", int(C_HIST), 31999990042, "CONCLUÍDO", 0, "mesma data, status diferentes"),
        ("R16", "SEM NUMERO", int(C_SINGLE), "PEDIDO", "CONCLUÍDO", 0, "acesso em texto"),
        ("R17", "INVALIDO", 12345678000100, 81999990001, "CONCLUÍDO", 0, "cnpj inválido"),
        ("R18", "DESLOCADA", int(C_MIXED), 85999990026, None, "CONCLUÍDO", "fila em F (deslocada)"),
    ]
    for row in rows:
        tlv.append(list(row))
    tlv["G2"] = "=1+1"  # fórmula: deve sair como valor calculado/texto, nunca quebrar

    cstv = wb.create_sheet("CSTV")
    cstv.append(CSTV_HEADERS)  # B=Acesso, C=Temporário, E=CNPJ/CPF, sem Fila Atual
    cstv.append([1, 81999990001, 81988880001, "ALFA", int(C_SINGLE), "Portada"])        # mesmo CNPJ+nº da TLV
    cstv.append([2, 85999990023, 85988880023, "GAMA", int(C_MIXED), "Pendente"])
    cstv.append([3, 11999990005, None, "DELTA", int(C_PAGED), "Portada"])                # "Portado": ignorado

    hidden = wb.create_sheet("Oculta")
    hidden.sheet_state = "hidden"
    hidden.append(["CNPJ", "Acesso", "fila atual", "x", "y"])
    hidden.append([int(C_SINGLE), 81999990001, "X", 1, 2])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


MAPPINGS = {
    "TLV": {"header_row": 1, "cnpj": "C", "acesso": "D", "fila_atual": "E", "numero_temporario": None,
            "expected_headers": {"C": "CNPJ", "D": "Acesso", "E": "fila atual"}},
    "CSTV": {"header_row": 1, "cnpj": "E", "acesso": "B", "fila_atual": None, "numero_temporario": "C",
             "expected_headers": {"E": "CNPJ/CPF", "B": "Acesso", "C": "Temporário"}},
}
