"""Execução completa: importação -> consultas -> exportação dos 3 arquivos + relatório; retomada."""
from __future__ import annotations

import threading

import pytest
from openpyxl import load_workbook

from core import paths
from core.engine import EngineState
from core.models import RunStatus, SearchResult
from phoenix import jobs
from phoenix.config import PhoenixSettings, SheetMapping
from phoenix.workbook import file_sha256
from tests.conftest import FakeAdapter, wait_until
from tests.phoenix.fixtures import C_MIXED, C_NONE, C_SINGLE, DATA, MAPPINGS, build_workbook, client_result


def mappings(*names) -> dict[str, SheetMapping]:
    return {n: SheetMapping(sheet=n, **MAPPINGS[n]) for n in (names or MAPPINGS)}


def settings(tmp_path, **kw) -> PhoenixSettings:
    s = PhoenixSettings(interval_seconds=0, timeout_seconds=15, max_retries=1, backoff_base_seconds=0.05,
                        output_dir=str(tmp_path / "saida"), autosave_every=1000, persist_session=False, headless=True)
    for k, v in kw.items():
        setattr(s, k, v)
    return s


class FakePhoenix(FakeAdapter):
    """Responde como o PhoenixAdapter, a partir dos dados fictícios, sem navegador."""

    id = "phoenix2"

    def __init__(self, delay: float = 0.0):
        super().__init__(delay=delay)

    def search(self, identifier):
        super().search(identifier)
        result = client_result(identifier)
        return SearchResult.found_with(result) if result else SearchResult.not_found("Nenhuma Solicitação foi encontrada.")


def engine_with(db, run_id, adapter, **kw):
    engine = jobs.make_engine(db, run_id, interactive=kw.pop("interactive", False), **kw)
    engine.adapter_factory = lambda s: adapter
    return engine


def rows(path, sheet):
    data = list(load_workbook(path)[sheet].iter_rows(values_only=True))
    return data[0], data[1:]


def test_full_job_exports_three_files_and_keeps_source(tmp_home, tmp_path):
    source = build_workbook(tmp_path / "entrada.xlsx")
    before = file_sha256(source)
    db = jobs.open_database()
    run_id, resumed = jobs.create_job(db, source, mappings(), settings(tmp_path))
    assert not resumed
    adapter = FakePhoenix()
    assert engine_with(db, run_id, adapter).run() == RunStatus.COMPLETED
    assert len(adapter.called()) == len(set(adapter.called())) == 7  # cada CNPJ distinto pesquisado 1 vez

    files = jobs.export(db, run_id)
    assert set(files) == {"EM_ANDAMENTO", "FALHA_PARCIAL_PORTADA", "EM_CONFLITO", "RELATORIO_TECNICO"}
    assert all(p.exists() and p.parent == tmp_path / "saida" for p in files.values())
    stamp = files["EM_ANDAMENTO"].stem.removeprefix("EM_ANDAMENTO_")
    assert len(stamp) == 15 and all(p.stem.endswith(stamp) for p in files.values())  # data e hora no nome

    header, data = rows(files["EM_ANDAMENTO"], "TLV")
    assert header[:7] == ("RADAR", "RAZÃO SOCIAL", "CNPJ", "Acesso", "fila atual", "fila atual", "OBS")
    assert header[7:] and header[7] == "[ROBÔ] Aba original"
    by_line = {r[header.index("[ROBÔ] Linha original")]: r for r in data}
    assert set(by_line) == {2, 6, 13, 15}
    assert by_line[2][4] == "Em andamento" and by_line[2][5] == 0      # só a "fila atual" mapeada (E) muda
    assert by_line[2][2] == int(C_SINGLE)
    # G2 é a fórmula "=1+1" gravada sem valor calculado: sai o valor (vazio), nunca a fórmula
    assert by_line[2][6] is None and by_line[6][6] == "formatado"

    _, cstv = rows(files["EM_ANDAMENTO"], "CSTV")
    cstv_header = load_workbook(files["EM_ANDAMENTO"])["CSTV"][1]
    assert [c.value for c in cstv_header][6] == "Fila Atual"           # criada só na exportação
    assert cstv[0][6] == "Em andamento"

    header, data = rows(files["EM_CONFLITO"], "TLV")
    fila = {r[header.index("[ROBÔ] Linha original")]: (r[4], r[header.index("[ROBÔ] Revisão")]) for r in data}
    assert fila[8] == ("Informação disponibilizada pelo sistema.", None)
    assert fila[9][0] == "Em conflito" and "motivo pendente" in fila[9][1]
    assert fila[4][0] == "Linha com débito na doadora"

    header, data = rows(files["FALHA_PARCIAL_PORTADA"], "TLV")
    assert {r[header.index("[ROBÔ] Linha original")] for r in data} == {5, 7}

    report = load_workbook(files["RELATORIO_TECNICO"])
    assert {"Resumo", "Revisar", "Sem correspondência", "CNPJ sem solicitações", "Erros de consulta",
            "Bilhetes sem linha", "Consultas", "Mapeamento"} <= set(report.sheetnames)
    sem = [r for r in report["Sem correspondência"].iter_rows(values_only=True)][1:]
    assert [r[1] for r in sem] == [11]
    assert C_NONE in [r[0] for r in report["CNPJ sem solicitações"].iter_rows(values_only=True)]

    assert file_sha256(source) == before                                # origem intacta


def test_only_classified_rows_go_to_each_file(tmp_home, tmp_path):
    source = build_workbook(tmp_path / "entrada.xlsx")
    db = jobs.open_database()
    run_id, _ = jobs.create_job(db, source, mappings("CSTV"), settings(tmp_path))
    engine_with(db, run_id, FakePhoenix()).run()
    files = jobs.export(db, run_id)
    assert load_workbook(files["FALHA_PARCIAL_PORTADA"]).sheetnames == ["Sem registros"]
    _, conflito = rows(files["EM_CONFLITO"], "CSTV")
    assert [r[1] for r in conflito] == [85999990023]


def test_pause_cancel_and_resume_after_interruption(tmp_home, tmp_path):    # cenário 16
    source = build_workbook(tmp_path / "entrada.xlsx")
    db = jobs.open_database()
    run_id, _ = jobs.create_job(db, source, mappings(), settings(tmp_path))
    adapter = FakePhoenix(delay=0.15)
    engine = engine_with(db, run_id, adapter, interactive=True)
    engine.start()
    assert wait_until(lambda: len(adapter.calls) >= 2)
    engine.pause()
    assert wait_until(lambda: engine.state == EngineState.PAUSED)
    frozen = len(adapter.calls)
    threading.Event().wait(0.3)
    assert len(adapter.calls) == frozen
    engine.stop()  # cancelar com segurança
    assert engine.join(10) and engine.final_status == RunStatus.STOPPED
    done = list(adapter.called())

    again, resumed = jobs.create_job(db, source, mappings(), settings(tmp_path))
    assert (again, resumed) == (run_id, True)                       # oferece retomar a mesma execução
    rest = FakePhoenix()
    assert engine_with(db, run_id, rest).run() == RunStatus.COMPLETED
    assert not set(done) & set(rest.called())                        # não repete CNPJs concluídos
    assert len(done) + len(rest.called()) == 7


@pytest.mark.browser
def test_end_to_end_with_browser_and_session_expiry(tmp_home, tmp_path):    # cenários 17 e 19 com navegador
    from core.browser import BrowserManager
    from phoenix_simulator.server import PhoenixSimulator
    from tests.phoenix.test_client import login

    source = build_workbook(tmp_path / "entrada.xlsx")
    db = jobs.open_database()
    with PhoenixSimulator(DATA) as sim, BrowserManager(headless=True, channel="msedge", timeout_seconds=15) as browser:
        run_id, _ = jobs.create_job(db, source, mappings("TLV"), settings(tmp_path, base_url=sim.base_url))
        login(browser.page, sim.base_url)
        saved = dict(sim.state.sessions)
        searches = []

        def on_event(event):
            if event.kind == "item_finished":
                searches.append(event.data["status"])
                if len(searches) == 2:
                    sim.state.expire_all()  # sessão expira no meio do lote

        engine = jobs.make_engine(db, run_id, interactive=False, on_event=on_event, browser=browser)
        assert engine.run() == RunStatus.NEEDS_OPERATOR                # modo sem operador: para e avisa
        sim.state.sessions.update(saved)                                # operador entrou de novo
        engine = jobs.make_engine(db, run_id, interactive=False, browser=browser)
        assert engine.run() == RunStatus.COMPLETED
        assert browser.browser.is_connected()                           # navegador emprestado continua aberto

    files = jobs.export(db, run_id)
    header, data = rows(files["EM_ANDAMENTO"], "TLV")
    assert {r[header.index("[ROBÔ] Linha original")] for r in data} == {2, 6, 13, 15}
    report = jobs.report(db, run_id)
    assert report.summary()["em_conflito"] == 4 and report.summary()["erro_consulta"] == 0
