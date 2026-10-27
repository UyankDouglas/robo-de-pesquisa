"""Normalização, leitura da planilha, correspondência CNPJ + número e regras de classificação."""
from __future__ import annotations

import pytest

from core.models import ConfigError, ItemStatus
from phoenix.config import MatchRules, SheetMapping, StatusRules
from phoenix.matcher import Category, CnpjQuery, Outcome, match
from phoenix.normalize import PhoneRules, normalize_cnpj, normalize_phone
from phoenix.workbook import detect_mappings, list_sheets, load_sheets
from tests.phoenix.fixtures import (C_HIST, C_MIXED, C_MULTI, C_NONE, C_PAGED, C_SINGLE, C_ZEROS, MAPPINGS,
                                    build_workbook, client_result)


def mappings(**changes) -> dict[str, SheetMapping]:
    out = {name: SheetMapping(sheet=name, **m) for name, m in MAPPINGS.items()}
    for name, attrs in changes.items():
        for key, value in attrs.items():
            setattr(out[name], key, value)
    return out


def all_queries(cnpjs) -> dict[str, CnpjQuery]:
    out = {}
    for c in cnpjs:
        result = client_result(c)
        out[c] = CnpjQuery(ItemStatus.FOUND, result) if result else CnpjQuery(ItemStatus.NOT_FOUND)
    return out


@pytest.fixture
def sheets(tmp_path):
    return load_sheets(build_workbook(tmp_path / "entrada.xlsx"), mappings())


@pytest.fixture
def report(sheets):
    rows = [r for s in sheets for r in s.rows]
    queries = all_queries({r.cnpj.value for r in rows if r.cnpj.value})
    return match(rows, queries, StatusRules(), MatchRules())


def by_row(report, sheet, row):
    return next(r for r in report.rows if r.row.sheet == sheet and r.row.row == row)


# ------------------------------------------------------------------ normalização
def test_cnpj_normalization_keeps_leading_zeros():
    assert normalize_cnpj(int(C_ZEROS)).value == C_ZEROS          # número do Excel sem zeros
    assert normalize_cnpj(float(int(C_ZEROS))).value == C_ZEROS    # float sem notação científica
    assert normalize_cnpj("11.222.333/0001-81").value == "11222333000181"
    assert normalize_cnpj("12345678000100").status == "invalido"
    assert normalize_cnpj("52998224725").status == "cpf"
    assert normalize_cnpj("12ABC34501DE35").status == "valido"     # formato alfanumérico novo


def test_phone_normalization_rules():
    assert normalize_phone("(81) 99999-0001") == "81999990001"
    assert normalize_phone(5581999990001) == "81999990001"         # DDI 55 removido (13 dígitos)
    assert normalize_phone(55999990001) == "55999990001"           # DDD 55 não é DDI
    assert normalize_phone(5581999990001, PhoneRules(strip_ddi_55=False)) == ""
    assert normalize_phone("PEDIDO") == "" and normalize_phone("8199999") == ""


# ------------------------------------------------------------------ planilha
def test_sheets_listing_hidden_and_mapping_detection(tmp_path):
    path = build_workbook(tmp_path / "e.xlsx")
    sheets = {s.name: s for s in list_sheets(path)}
    assert sheets["Oculta"].hidden and not sheets["TLV"].hidden
    detected = detect_mappings(path)
    # "fila atual" duplicada (E texto, F numérico): a detecção escolhe pela validação do conteúdo
    assert detected["TLV"].fila_atual == "E" and "repetido" in detected["TLV"].note
    assert (detected["CSTV"].cnpj, detected["CSTV"].acesso, detected["CSTV"].fila_atual) == ("E", "B", None)


def test_mapping_header_mismatch_is_refused(tmp_path):
    path = build_workbook(tmp_path / "e.xlsx")
    wrong = mappings(TLV={"fila_atual": "F", "expected_headers": {"C": "CNPJ", "D": "Acesso", "F": "STATUS"}})
    with pytest.raises(ConfigError, match="era esperado"):
        load_sheets(path, wrong)


def test_number_in_date_formatted_column_is_preserved(tmp_path):
    """Planilha real: nº de contrato numa coluna com formato de data não pode virar '#VALUE!'."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "TLV"
    ws.append(["RADAR", "RAZÃO SOCIAL", "CNPJ", "Acesso", "fila atual", "Contrato", "Ativação"])
    ws.append(["R1", "X", 10000000000145, 81999990001, "OK", 2026671220003, 45992])
    ws["F2"].number_format = "dd/mm/yyyy"   # número enorme formatado como data
    ws["G2"].number_format = "dd/mm/yyyy"   # data válida continua data
    path = tmp_path / "datas.xlsx"
    wb.save(path)
    mapping = SheetMapping(sheet="TLV", cnpj="C", acesso="D", fila_atual="E")
    row = load_sheets(path, {"TLV": mapping})[0].rows[0]
    assert row.values[5] == 2026671220003
    assert row.values[6].year == 2025


def test_shifted_fila_atual_row_is_flagged(sheets):
    tlv = sheets[0]
    shifted = next(r for r in tlv.rows if r.row == 19)
    assert shifted.warnings and "deslocamento" in shifted.warnings[0]
    assert any("deslocamento" in issue for issue in tlv.issues)


# ------------------------------------------------------------------ correspondência
def test_single_request_in_progress(report):                       # cenários 1 e 5
    res = by_row(report, "TLV", 2)
    assert (res.outcome, res.category, res.fila_nova) == (Outcome.CLASSIFICADO, Category.EM_ANDAMENTO, "Em andamento")


def test_multiple_requests_use_ticket_status_not_request_status(report):   # cenários 2 e 4
    assert by_row(report, "TLV", 3).outcome == Outcome.OUTRO_STATUS  # bilhete "Portado" (solicitação "Em andamento")
    conflict = by_row(report, "TLV", 4)
    assert conflict.category == Category.EM_CONFLITO and conflict.fila_nova == "Linha com débito na doadora"
    assert by_row(report, "TLV", 5).category == Category.FALHA_PARCIAL_PORTADA


def test_one_request_many_tickets(report):                          # cenários 3, 6, 7, 8 e 9
    assert by_row(report, "TLV", 6).category == Category.EM_ANDAMENTO        # acesso formatado "(85) 99999-0021"
    falha = by_row(report, "TLV", 7)
    assert falha.category == Category.FALHA_PARCIAL_PORTADA and falha.fila_nova == "Falha parcial / portada"
    conflito = by_row(report, "TLV", 8)
    assert conflito.fila_nova == "Informação disponibilizada pelo sistema." and not conflito.revisao
    sem_motivo = by_row(report, "TLV", 9)
    assert sem_motivo.category == Category.EM_CONFLITO and sem_motivo.fila_nova == "Em conflito"
    assert "motivo pendente" in " ".join(sem_motivo.revisao)
    assert by_row(report, "TLV", 10).outcome == Outcome.OUTRO_STATUS          # "Cancelado"


def test_not_found_and_mismatched_number(report):                   # cenários 10 e 11
    assert by_row(report, "TLV", 12).outcome == Outcome.CNPJ_SEM_SOLICITACOES
    assert by_row(report, "TLV", 11).outcome == Outcome.SEM_CORRESPONDENCIA
    unmatched = {t.numero for t in report.unmatched_tickets}
    assert "85999990099" not in unmatched and "11999990001" in unmatched  # bilhetes sem linha vão ao relatório


def test_repeated_cnpj_and_cross_sheet_rows(report):                # cenário 12
    alfa = [r for r in report.rows if r.row.cnpj.value == C_SINGLE]
    assert {(r.row.sheet, r.row.row, r.outcome) for r in alfa} == {
        ("TLV", 2, Outcome.CLASSIFICADO), ("TLV", 17, Outcome.SEM_NUMERO), ("CSTV", 2, Outcome.CLASSIFICADO)}
    classified = [r for r in report.rows if r.ticket]
    assert classified and all(r.ticket.numero == r.row.acesso and r.ticket.cnpj == r.row.cnpj.value
                              for r in classified)  # cada linha com o bilhete do SEU CNPJ + número
    cross = [r for r in classified if r.row.acesso == "85999990023"]
    assert {r.row.sheet for r in cross} == {"TLV", "CSTV"}  # mesma linha telefônica em duas abas


def test_leading_zeros_ddi_pagination_and_history(report):         # cenários 13 e 17
    zeros = by_row(report, "TLV", 14)
    assert zeros.row.cnpj.value == C_ZEROS and zeros.fila_nova == "Titular divergente"
    assert by_row(report, "TLV", 13).category == Category.EM_ANDAMENTO      # bilhete da 3ª página
    assert by_row(report, "CSTV", 4).outcome == Outcome.OUTRO_STATUS
    hist = by_row(report, "TLV", 15)
    assert hist.category == Category.EM_ANDAMENTO and hist.ticket.solicitacao["Data Solicitação"] == "20/08/2026 10:00"
    same_date = by_row(report, "TLV", 16)
    assert same_date.outcome == Outcome.REVISAR and "mesma data" in " ".join(same_date.revisao)


def test_history_rule_review_only(sheets):
    rows = [r for s in sheets for r in s.rows if r.cnpj.value == C_HIST]
    rep = match(rows, all_queries([C_HIST]), StatusRules(), MatchRules(multiplos_bilhetes="revisar"))
    assert all(r.outcome == Outcome.REVISAR for r in rep.rows if r.row.acesso)


def test_invalid_cnpj_and_text_access(report):
    assert by_row(report, "TLV", 18).outcome == Outcome.CNPJ_INVALIDO
    assert by_row(report, "TLV", 17).outcome == Outcome.SEM_NUMERO


def test_portado_is_not_falha_parcial_unless_configured(sheets):
    rows = [r for s in sheets for r in s.rows if r.cnpj.value == C_MULTI]
    default = match(rows, all_queries([C_MULTI]), StatusRules(), MatchRules())
    assert by_row(default, "TLV", 3).outcome == Outcome.OUTRO_STATUS
    configured = StatusRules(falha_parcial_portada=["Falha parcial", "Portado"])
    rep = match(rows, all_queries([C_MULTI]), configured, MatchRules())
    assert by_row(rep, "TLV", 3).category == Category.FALHA_PARCIAL_PORTADA   # só se o operador configurar


def test_temporary_number_only_when_enabled_and_unambiguous(sheets):
    rows = [r for s in sheets for r in s.rows if r.sheet == "CSTV" and r.cnpj.value == C_SINGLE]
    rows[0].acesso = "81000000000"  # acesso não bate; temporário (81988880001) também não existe no bilhete
    rep = match(rows, all_queries([C_SINGLE]), StatusRules(), MatchRules(usar_numero_temporario=True))
    assert rep.rows[0].outcome == Outcome.SEM_CORRESPONDENCIA


def test_pending_and_error_queries(sheets):
    rows = [r for s in sheets for r in s.rows if r.cnpj.value in (C_SINGLE, C_PAGED)]
    queries = {C_SINGLE: CnpjQuery(ItemStatus.PENDING), C_PAGED: CnpjQuery(ItemStatus.ERROR, None, "timeout")}
    rep = match(rows, queries, StatusRules(), MatchRules())
    assert {r.outcome for r in rep.rows if r.row.cnpj.value == C_PAGED} == {Outcome.ERRO_CONSULTA}
    assert Outcome.PENDENTE in {r.outcome for r in rep.rows if r.row.cnpj.value == C_SINGLE}
