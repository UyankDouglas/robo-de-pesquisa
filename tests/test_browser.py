"""Automação real do navegador (Playwright + Edge instalado) contra o site local de demonstração."""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from adapters.base_adapter import AdapterContext
from adapters.declarative_adapter import SiteDefinition, make_adapter_class
from adapters.demo_adapter import DemoBrowserAdapter, DemoLoginAdapter
from core import paths
from core.browser import BrowserManager
from core.engine import ProcessingEngine
from core.models import (
    HumanVerificationRequired,
    ItemStatus,
    RunStatus,
    SessionExpiredError,
    TransientError,
)
from core.session import SessionStore
from demo_site.server import DemoOptions, DemoServer
from tests.conftest import fast_settings, make_import

pytestmark = pytest.mark.browser
CHANNEL = "msedge"


def _open(adapter, browser, timeout=10):
    adapter.open(AdapterContext(timeout_seconds=timeout, browser=browser))
    return adapter


def test_demo_adapter_results(demo_server):
    with BrowserManager(headless=True, channel=CHANNEL, timeout_seconds=10) as browser:
        adapter = _open(DemoBrowserAdapter({"base_url": demo_server.url, "autostart_server": False}), browser)
        found = adapter.search("12345678000190")
        assert found.found
        assert found.data == {
            "nome": "Empresa Fictícia Alfa LTDA", "situacao": "ATIVA", "data_registro": "15/03/2010",
            "municipio": "São Paulo", "uf": "SP", "detalhes": "Comércio varejista de artigos de demonstração",
        }
        assert adapter.search("00012345000101").data["nome"] == "Gama Zeros à Esquerda ME"
        missing = adapter.search("98765432000110")
        assert not missing.found and "Nenhum resultado" in missing.message
        with pytest.raises(TransientError):
            adapter.search("ERRO-1")
        with pytest.raises(HumanVerificationRequired):
            adapter.search("CAPTCHA-1")
        assert adapter.search("PROTOCOLO-2024-0001").found  # não só CNPJ


def test_timeout_is_classified_as_transient():
    with DemoServer("127.0.0.1", 0, DemoOptions(latency=0, slow_seconds=4)) as server, \
            BrowserManager(headless=True, channel=CHANNEL, timeout_seconds=1.5) as browser:
        adapter = _open(DemoBrowserAdapter({"base_url": server.url, "autostart_server": False}), browser, timeout=1.5)
        with pytest.raises(Exception) as info:
            adapter.search("LENTO-1")
        assert adapter.classify_error(info.value) == "transient"
        adapter.recover(info.value)
        assert adapter.search("12345678000190").found  # segue funcionando depois do timeout


def _operator_logs_in(page) -> None:
    """Simula o operador digitando no navegador (no uso real é a pessoa que faz isso)."""
    page.get_by_test_id("login-user").fill("demo")
    page.get_by_test_id("login-password").fill("demo123")
    page.get_by_test_id("login-submit").click()
    page.wait_for_url("**/")


def test_assisted_login_persisted_session_and_expiry(login_server, tmp_home):
    options = {"base_url": login_server.url, "autostart_server": False}
    store = SessionStore()
    with BrowserManager(headless=True, channel=CHANNEL, timeout_seconds=10) as browser:
        adapter = _open(DemoLoginAdapter(options), browser)
        assert adapter.requires_login and not adapter.is_logged_in()
        with pytest.raises(SessionExpiredError):
            adapter.search("12345678000190")
        assert "Faça login" in adapter.begin_login()
        _operator_logs_in(browser.page)
        assert adapter.is_logged_in()
        assert adapter.search("12345678000190").found
        assert store.save(adapter.id, adapter.storage_state())

    # nova abertura do navegador reaproveita a sessão protegida, sem novo login
    with BrowserManager(headless=True, channel=CHANNEL, timeout_seconds=10,
                        storage_state=store.load("demo_login")) as browser:
        adapter = _open(DemoLoginAdapter(options), browser)
        assert adapter.is_logged_in()
        login_server.state.sessions.clear()  # o site expira a sessão
        with pytest.raises(SessionExpiredError):
            adapter.search("11222333000144")
        assert not adapter.is_logged_in()


def test_declarative_json_adapter(demo_server):
    raw = json.loads((paths.BUNDLED_SITES_DIR / "demo_declarativo.json").read_text(encoding="utf-8"))
    cls = make_adapter_class(SiteDefinition.from_dict(raw, "demo_declarativo.json"))
    with BrowserManager(headless=True, channel=CHANNEL, timeout_seconds=10) as browser:
        adapter = _open(cls({"search_url": demo_server.url + "/", "autostart_server": False}), browser)
        result = adapter.search("11222333000144")
        assert result.found and result.data["nome"] == "Beta Serviços de Teste S.A." and result.data["uf"] == "MG"
        assert set(result.data) == {"nome", "situacao", "data_registro", "municipio", "uf"}
        assert not adapter.search("NF-123").found
        with pytest.raises(HumanVerificationRequired):
            adapter.search("CAPTCHA-2")
        with pytest.raises(TransientError):
            adapter.search("ERRO-2")


def test_engine_end_to_end_with_real_browser(demo_server, db, repo, tmp_path):
    ids = ["12345678000190", "98765432000110", "INSTAVEL-77", "ERRO-9", "00012345000101", "12345678000190"]
    settings = fast_settings(
        adapter_id="demo", adapter_options={"base_url": demo_server.url, "autostart_server": False},
        headless=True, browser_channel=CHANNEL, timeout_seconds=10, max_retries=2, backoff_base_seconds=0.1,
        output_dir=str(tmp_path / "saida"),
    )
    run_id = repo.create_run("e2e", make_import(tmp_path, ids), settings)
    engine = ProcessingEngine(db, run_id, interactive=False)
    assert engine.run() == RunStatus.COMPLETED

    status = {i.identifier: i.status for i in repo.items(run_id) if i.status != ItemStatus.DUPLICATE}
    assert status == {
        "12345678000190": ItemStatus.FOUND,
        "98765432000110": ItemStatus.NOT_FOUND,
        "INSTAVEL-77": ItemStatus.FOUND,      # recuperou após 2 falhas temporárias
        "ERRO-9": ItemStatus.ERROR,
        "00012345000101": ItemStatus.FOUND,
    }
    assert engine.output_path.parent == tmp_path / "saida"
    rows = list(load_workbook(engine.output_path)["Resultados"].iter_rows(values_only=True))
    header = rows[0]
    by_doc = {r[header.index("DOC")]: r for r in rows[1:]}
    assert by_doc["00012345000101"][header.index("NOME / RAZÃO SOCIAL")] == "Gama Zeros à Esquerda ME"
    assert by_doc["ERRO-9"][header.index("SITUAÇÃO")] == "Erro"
    assert by_doc["ERRO-9"][header.index("TENTATIVAS")] == 3
    assert len(rows) == 7  # cabeçalho + 6 linhas originais (inclui o duplicado)
