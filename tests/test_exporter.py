"""Exportação: dados originais preservados, filtros, duplicados, gravação segura e arquivo válido."""
from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook

from core.models import ItemStatus
from excel import exporter
from excel.exporter import build_dataframe, export_run
from tests.conftest import fast_settings, make_import, make_workbook
from excel.importer import ImportResult, build_processing_list, read_sheet

LABELS = [("nome", "NOME"), ("situacao", "SITUAÇÃO")]


def _run_with_results(repo, tmp_path) -> int:
    imported = make_import(tmp_path, ["00012345000101", "98765432000110", "ERRO-1", "00012345000101", "", "=1+1"])
    run_id = repo.create_run("teste", imported, fast_settings())
    items = {i.identifier: i for i in repo.items(run_id) if i.status == ItemStatus.PENDING}
    repo.complete_item(items["00012345000101"].id, ItemStatus.FOUND, {"nome": "Empresa X", "situacao": "ATIVA"})
    repo.complete_item(items["98765432000110"].id, ItemStatus.NOT_FOUND, None, "Nenhum resultado")
    repo.complete_item(items["ERRO-1"].id, ItemStatus.ERROR, None, "TransientError: HTTP 500")
    return run_id


def _sheet_rows(path: Path, sheet: str = "Resultados") -> list[dict]:
    wb = load_workbook(path)
    rows = list(wb[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_export_all_preserves_original_and_adds_results(repo, tmp_path):
    run_id = _run_with_results(repo, tmp_path)
    out = export_run(repo, run_id, tmp_path / "out" / "resultado.xlsx", "all", LABELS)
    rows = _sheet_rows(out)
    assert len(rows) == 6  # todas as linhas originais, inclusive duplicado e vazio
    first = rows[0]
    assert first["DOC"] == "00012345000101"  # zeros à esquerda preservados (texto)
    assert first["ID"] == 1                 # demais colunas numéricas continuam números
    assert first["NOME"] == "Empresa X"
    # coluna do site "SITUAÇÃO" colide com a de controle "SITUAÇÃO": ambas existem sem sobrescrever
    assert first["SITUAÇÃO"] == "ATIVA" and first["SITUAÇÃO (consulta)"] == "Concluído"
    assert first["RESULTADO"] == "Informação encontrada"
    assert first["DATA/HORA DA CONSULTA"] is not None
    assert rows[1]["SITUAÇÃO (consulta)"] == "Sem resultado" and rows[1]["RESULTADO"] == "Não encontrado"
    assert rows[2]["SITUAÇÃO (consulta)"] == "Erro" and "HTTP 500" in rows[2]["MENSAGEM"]
    assert rows[3]["SITUAÇÃO (consulta)"] == "Duplicado" and rows[3]["NOME"] == "Empresa X"
    assert rows[4]["SITUAÇÃO (consulta)"] == "Inválido"
    assert rows[5]["DOC"] == "=1+1"  # texto, nunca fórmula
    wb = load_workbook(out)
    assert wb["Resultados"]["B7"].data_type == "s"
    assert "Resumo" in wb.sheetnames


def test_export_filters(repo, tmp_path):
    run_id = _run_with_results(repo, tmp_path)
    by_filter = {f: build_dataframe(repo, run_id, f, LABELS) for f in ("found", "not_found", "errors", "processed", "pending")}
    assert list(by_filter["found"]["DOC"]) == ["00012345000101"]
    assert list(by_filter["not_found"]["DOC"]) == ["98765432000110"]
    assert list(by_filter["errors"]["DOC"]) == ["ERRO-1"]
    assert len(by_filter["processed"]) == 2
    assert list(by_filter["pending"]["DOC"]) == ["=1+1"]
    # nenhuma exportação filtrada repete identificadores
    for frame in by_filter.values():
        assert not frame["DOC"].duplicated().any()


def test_atomic_save_leaves_no_temp_and_replaces(repo, tmp_path):
    run_id = _run_with_results(repo, tmp_path)
    target = tmp_path / "saida.xlsx"
    export_run(repo, run_id, target)
    first_size = target.stat().st_size
    export_run(repo, run_id, target)  # regrava por cima sem erro
    assert target.stat().st_size > 0 and first_size > 0
    assert [p.name for p in tmp_path.iterdir() if p.name.startswith(".~")] == []


def test_locked_target_falls_back_to_new_file(repo, tmp_path, monkeypatch):
    run_id = _run_with_results(repo, tmp_path)
    target = tmp_path / "aberto_no_excel.xlsx"
    real_replace = os.replace

    def fake_replace(src, dst):
        if Path(dst) == target:
            raise PermissionError("arquivo em uso")
        return real_replace(src, dst)

    monkeypatch.setattr(exporter.os, "replace", fake_replace)
    monkeypatch.setattr(exporter.time, "sleep", lambda s: None)
    saved = export_run(repo, run_id, target)
    assert saved != target and saved.name.startswith("aberto_no_excel_") and saved.exists()
    load_workbook(saved)  # arquivo válido


def test_failed_write_does_not_corrupt_existing_file(repo, tmp_path, monkeypatch):
    run_id = _run_with_results(repo, tmp_path)
    target = tmp_path / "estavel.xlsx"
    export_run(repo, run_id, target)
    before = target.read_bytes()

    def broken(*args, **kwargs):
        raise RuntimeError("queda no meio da gravação")

    monkeypatch.setattr(exporter, "_write_workbook", broken)
    try:
        export_run(repo, run_id, target)
    except RuntimeError:
        pass
    assert target.read_bytes() == before  # o arquivo anterior continua íntegro
    assert not any(p.name.startswith(".~") for p in tmp_path.iterdir())


def test_illegal_characters_are_removed(repo, tmp_path):
    path = make_workbook(tmp_path / "ctrl.xlsx", [["DOC"], ["ABC"]])
    frame, info, _ = read_sheet(path)
    imported = ImportResult(str(path), info.name, "DOC", info.columns, 1, build_processing_list(frame, "DOC"))
    run_id = repo.create_run("x", imported, fast_settings())
    item = repo.items(run_id)[0]
    repo.complete_item(item.id, ItemStatus.FOUND, {"nome": "texto\x00com\x07controle"})
    out = export_run(repo, run_id, tmp_path / "ok.xlsx", "all", LABELS)
    assert _sheet_rows(out)[0]["NOME"] == "textocomcontrole"
