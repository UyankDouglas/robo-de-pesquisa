"""Importação de Excel: abas/colunas, preservação de identificadores, duplicados, vazios e intervalo."""
from __future__ import annotations

import pytest

from core.config import ImportOptions
from core.models import ConfigError, ItemStatus
from excel.importer import build_processing_list, import_excel, inspect_workbook, read_sheet
from tests.conftest import make_workbook


def test_inspect_lists_sheets_and_columns(tmp_path):
    path = make_workbook(tmp_path / "a.xlsx", [["ID", "DOCUMENTO", "STATUS"], [1, "123", "Pendente"]])
    from openpyxl import load_workbook

    wb = load_workbook(path)
    extra = wb.create_sheet("Outra")
    extra.append(["CODIGO", "NOME"])
    extra.append(["A1", "Fulano"])
    extra.append(["A2", "Beltrano"])
    wb.save(path)

    infos = inspect_workbook(path)
    assert [i.name for i in infos] == ["Dados", "Outra"]
    assert infos[0].columns == ["ID", "DOCUMENTO", "STATUS"]
    assert infos[0].data_rows == 1
    assert infos[1].columns == ["CODIGO", "NOME"]
    assert infos[1].data_rows == 2


def test_leading_zeros_preserved_in_text_cells(tmp_path):
    path = make_workbook(tmp_path / "z.xlsx", [["DOC"], ["00012345000101"], ["0001"], ["000"]])
    frame, _, _ = read_sheet(path)
    records = build_processing_list(frame, "DOC")
    assert [r.identifier for r in records] == ["00012345000101", "0001", "000"]


def test_numeric_cells_keep_identifiers(tmp_path):
    path = make_workbook(
        tmp_path / "n.xlsx",
        [["DOC"], [98765000155], [12345678000190], [42.0], [1.5]],
        formats={"A2": "00000000000000"},  # formato de zeros: o Excel exibe 00098765000155
    )
    frame, _, _ = read_sheet(path)
    ids = [r.identifier for r in build_processing_list(frame, "DOC")]
    assert ids == ["00098765000155", "12345678000190", "42", "1.5"]  # sem notação científica nem ".0"


def test_duplicates_are_marked_and_first_occurrence_kept(tmp_path):
    path = make_workbook(tmp_path / "d.xlsx", [["DOC"], ["111"], ["222"], ["111"], [" 222 "], ["333"]])
    frame, _, _ = read_sheet(path)
    records = build_processing_list(frame, "DOC")
    statuses = [(r.identifier, r.status) for r in records]
    assert statuses == [
        ("111", ItemStatus.PENDING),
        ("222", ItemStatus.PENDING),
        ("111", ItemStatus.DUPLICATE),
        ("222", ItemStatus.DUPLICATE),  # espaços extras não escondem duplicados
        ("333", ItemStatus.PENDING),
    ]
    assert records[2].duplicate_of_row == 2 and records[3].duplicate_of_row == 3


def test_empty_identifiers_and_blank_rows(tmp_path):
    rows = [["ID", "DOC"], [1, "AAA"], [2, None], [None, None], [3, "   "], [4, "BBB"], [None, None], [None, None]]
    path = make_workbook(tmp_path / "e.xlsx", rows)
    frame, info, blank = read_sheet(path)
    assert info.data_rows == 4 and blank == 1  # linha totalmente vazia no meio; vazias no fim são descartadas
    records = build_processing_list(frame, "DOC")
    assert [r.status for r in records] == [ItemStatus.PENDING, ItemStatus.INVALID, ItemStatus.INVALID, ItemStatus.PENDING]
    assert records[1].message == "Identificador vazio"
    assert [r.row_number for r in records] == [2, 3, 5, 6]  # linha real do Excel


def test_range_selection(tmp_path):
    path = make_workbook(tmp_path / "r.xlsx", [["DOC"]] + [[f"ID{i}"] for i in range(1, 11)])
    frame, _, _ = read_sheet(path)
    records = build_processing_list(frame, "DOC", start=3, end=5)
    assert [r.identifier for r in records] == ["ID3", "ID4", "ID5"]
    assert [r.position for r in records] == [3, 4, 5]
    assert len(build_processing_list(frame, "DOC", start=8)) == 3
    with pytest.raises(ConfigError):
        build_processing_list(frame, "DOC", start=50)


def test_header_detected_below_title_rows(tmp_path):
    rows = [["Relatório de pendências"], [None], ["CODIGO", "NOME"], ["X1", "Ana"], ["X2", "Bruno"]]
    path = make_workbook(tmp_path / "h.xlsx", rows)
    infos = inspect_workbook(path)
    assert infos[0].header_row == 1  # 1ª linha preenchida
    frame, info, _ = read_sheet(path, header_row=3)
    assert info.columns == ["CODIGO", "NOME"]
    assert [r.identifier for r in build_processing_list(frame, "NOME")] == ["Ana", "Bruno"]


def test_duplicate_and_empty_header_names(tmp_path):
    path = make_workbook(tmp_path / "c.xlsx", [["DOC", None, "doc"], ["1", "x", "2"]])
    assert inspect_workbook(path)[0].columns == ["DOC", "Coluna B", "doc (2)"]


def test_zero_pad_and_validation(tmp_path):
    path = make_workbook(tmp_path / "v.xlsx", [["CNPJ"], ["11222333000181"], [11222333000181], ["11222333000182"],
                                               ["191"], ["123"]])
    frame, _, _ = read_sheet(path)
    records = build_processing_list(frame, "CNPJ", zero_pad=14, validation="cnpj")
    assert [r.identifier for r in records] == [
        "11222333000181", "11222333000181", "11222333000182", "00000000000191", "00000000000123"]
    assert [r.status for r in records] == [
        ItemStatus.PENDING, ItemStatus.DUPLICATE, ItemStatus.INVALID, ItemStatus.PENDING, ItemStatus.INVALID]
    assert records[2].message == "CNPJ inválido"


def test_accepts_any_kind_of_identifier(tmp_path):
    path = make_workbook(tmp_path / "g.xlsx", [["CHAVE"], ["PROTOCOLO-2024-0001"], ["Maria da Silva"], ["ABC/123"]])
    result = import_excel(ImportOptions(file=str(path), column="CHAVE"))
    assert result.to_process == 3
    assert result.records[1].original == {"CHAVE": "Maria da Silva"}


def test_import_excel_summary_and_errors(tmp_path):
    path = make_workbook(tmp_path / "s.xlsx", [["ID", "DOC", "STATUS"], [1, "A", "P"], [2, "A", "P"], [3, None, "P"]])
    result = import_excel(ImportOptions(file=str(path), column="DOC"))
    assert (result.to_process, len(result.duplicates), len(result.invalid)) == (1, 1, 1)
    assert result.records[0].original == {"ID": "1", "DOC": "A", "STATUS": "P"}
    with pytest.raises(ConfigError, match="Coluna 'XYZ'"):
        import_excel(ImportOptions(file=str(path), column="XYZ"))
    with pytest.raises(ConfigError, match="Aba"):
        import_excel(ImportOptions(file=str(path), column="DOC", sheet="Nao existe"))
    with pytest.raises(ConfigError, match="não encontrada"):
        import_excel(ImportOptions(file=str(tmp_path / "nada.xlsx"), column="DOC"))
    bad = tmp_path / "corrompido.xlsx"
    bad.write_text("isto não é um xlsx")
    with pytest.raises(ConfigError, match="Não foi possível abrir"):
        import_excel(ImportOptions(file=str(bad), column="DOC"))
