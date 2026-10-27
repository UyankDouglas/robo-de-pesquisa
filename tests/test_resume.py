"""Retomada: recuperação após queda, sem repetir concluídos, repetição de erros e trava entre processos."""
from __future__ import annotations

import subprocess
import sys
import textwrap

from core.config import ImportOptions, JobConfig
from core.engine import ProcessingEngine
from core.jobs import mark_stale_runs, prepare_job_run
from core.locking import RunLock
from core.models import ItemStatus, RunStatus, SearchResult, TransientError
from database.repository import Repository
from tests.conftest import ROOT, FakeAdapter, fast_settings, make_import, make_workbook


def _run(db, run_id, adapter):
    return ProcessingEngine(db, run_id, adapter_factory=lambda s: adapter, interactive=False).run()


def test_recovery_after_crash_does_not_repeat_completed(repo, db, tmp_path):
    run_id = repo.create_run("t", make_import(tmp_path, ["A", "B", "C", "D"]), fast_settings())
    items = repo.items(run_id)
    # simula queda: A e B concluídos, C "em processamento", processo morto sem finalizar
    repo.complete_item(items[0].id, ItemStatus.FOUND, {"nome": "a"})
    repo.complete_item(items[1].id, ItemStatus.NOT_FOUND)
    repo.mark_processing(items[2].id)
    repo.set_run_status(run_id, RunStatus.RUNNING)

    assert mark_stale_runs(db) == 1  # ninguém segura a trava -> interrompida
    assert repo.get_run(run_id).status == RunStatus.INTERRUPTED

    adapter = FakeAdapter()
    assert _run(db, run_id, adapter) == RunStatus.COMPLETED
    assert adapter.called() == ["C", "D"]  # retoma do ponto exato
    assert repo.items(run_id)[0].result == {"nome": "a"}  # resultado antigo preservado


def test_completed_records_only_rerun_on_request(repo, db, tmp_path):
    run_id = repo.create_run("t", make_import(tmp_path, ["A", "B"]), fast_settings())
    _run(db, run_id, FakeAdapter())
    again = FakeAdapter()
    _run(db, run_id, again)
    assert again.called() == []  # nada a refazer
    assert repo.reset_items(run_id, [ItemStatus.FOUND]) == 2  # pedido explícito
    _run(db, run_id, again)
    assert again.called() == ["A", "B"]


def test_retry_only_errors(repo, db, tmp_path):
    run_id = repo.create_run("t", make_import(tmp_path, ["A", "B", "C"]), fast_settings(max_retries=0))
    _run(db, run_id, FakeAdapter({"B": [TransientError("falhou")]}))
    assert repo.stats(run_id).error == 1
    assert repo.reset_items(run_id, [ItemStatus.ERROR]) == 1
    retry = FakeAdapter()
    _run(db, run_id, retry)
    assert retry.called() == ["B"]
    item_b = repo.items(run_id)[1]
    assert item_b.status == ItemStatus.FOUND and item_b.attempts == 2  # histórico de tentativas acumulado
    assert [a.outcome for a in repo.attempts(item_b.id)] == ["error", "found"]


def test_job_resume_reuses_incomplete_run(db, tmp_path):
    path = make_workbook(tmp_path / "job.xlsx", [["DOC"], ["A"], ["B"], ["C"]])
    job = JobConfig("job", ImportOptions(file=str(path), column="DOC"), fast_settings())
    run_id, resumed = prepare_job_run(db, job)
    assert not resumed
    repo = Repository(db.connect())
    first = repo.items(run_id)[0]
    repo.complete_item(first.id, ItemStatus.FOUND, {})
    repo.close()
    again, resumed = prepare_job_run(db, job)
    assert (again, resumed) == (run_id, True)
    adapter = FakeAdapter()
    _run(db, run_id, adapter)
    assert adapter.called() == ["B", "C"]
    third, resumed = prepare_job_run(db, job)  # tudo concluído -> nova execução
    assert third != run_id and not resumed


def test_lock_is_released_when_process_dies(tmp_home):
    from core import paths

    code = textwrap.dedent(f"""
        import sys, time
        sys.path.insert(0, {str(ROOT)!r})
        from pathlib import Path
        from core.locking import RunLock
        RunLock(77, Path({str(paths.LOCKS_DIR)!r})).acquire()
        print("locked", flush=True)
        time.sleep(60)
    """)
    proc = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.PIPE, text=True)
    try:
        assert proc.stdout.readline().strip() == "locked"
        assert RunLock.is_locked(77)  # outra instância detecta
    finally:
        proc.kill()
        proc.wait(10)
    assert not RunLock.is_locked(77)  # o SO libera a trava do processo morto


def test_interrupted_run_keeps_partial_excel(repo, db, tmp_path):
    run_id = repo.create_run("t", make_import(tmp_path, ["A", "B", "C"]), fast_settings(autosave_every=1))
    adapter = FakeAdapter({"B": [SearchResult.not_found()]})
    engine = ProcessingEngine(db, run_id, adapter_factory=lambda s: adapter, interactive=False)

    def stop_after_b(identifier):
        if identifier == "B":
            engine.stop()

    adapter.on_search = stop_after_b
    assert engine.run() == RunStatus.STOPPED
    from openpyxl import load_workbook

    rows = list(load_workbook(engine.output_path)["Resultados"].iter_rows(values_only=True))
    status_col = rows[0].index("SITUAÇÃO")
    assert [r[status_col] for r in rows[1:]] == ["Concluído", "Sem resultado", "Pendente"]
