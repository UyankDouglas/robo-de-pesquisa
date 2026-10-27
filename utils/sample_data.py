"""Gera a planilha de exemplo usada na demonstração e nos testes manuais."""
from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def create_sample_workbook(path: Path, lote_rows: int = 200, seed: int = 42) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Registros"
    ws.append(["ID", "DOCUMENTO", "STATUS"])
    rows = [
        (1, "12345678000190", "Pendente"),   # encontrado
        (2, "98765432000110", "Pendente"),   # sem resultado
        (3, "11222333000144", "Pendente"),   # encontrado
        (4, "00012345000101", "Pendente"),   # zeros à esquerda (texto)
        (5, "PROTOCOLO-2024-0001", "Pendente"),
        (6, "12345678000190", "Pendente"),   # duplicado da linha 2
        (7, None, "Pendente"),               # identificador vazio
        (8, "INSTAVEL-001", "Pendente"),     # falha temporária que se resolve
        (9, "ERRO-500", "Pendente"),         # erro persistente
        (10, "NF-00001", "Pendente"),        # sem resultado
    ]
    for row in rows:
        ws.append(list(row))
    # identificador gravado como NÚMERO com formato de zeros (comum em planilhas reais)
    ws.append([11, 98765000155, "Pendente"])
    ws.cell(row=ws.max_row, column=2).number_format = "00000000000000"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["B"].width = 24

    lote = wb.create_sheet("Lote")
    lote.append(["SEQ", "CODIGO", "OBSERVACAO"])
    rnd = random.Random(seed)
    for i in range(1, lote_rows + 1):
        lote.append([i, f"{rnd.randrange(10**13):014d}", f"Registro de volume {i}"])
    for cell in lote[1]:
        cell.font = Font(bold=True)
    lote.column_dimensions["B"].width = 20

    outros = wb.create_sheet("Protocolos e nomes")
    outros.append(["TIPO", "IDENTIFICADOR"])
    for row in [("Protocolo", "PROTOCOLO-2024-0001"), ("Código", "ABC-123"), ("Nome", "Maria Exemplo"),
                ("Nome", "José da Silva Fictício"), ("Código", "NF-INEXISTENTE")]:
        outros.append(list(row))
    for cell in outros[1]:
        cell.font = Font(bold=True)

    wb.save(path)
    return path
